"""
export_channels_sheet.py
-------------------------
Tạo file Excel (.xlsx) và CSV (.csv) thể hiện toàn bộ bảng tra cứu:
- Thư mục đầu vào (Input)
- Thư mục đầu ra (Output)
- Link kênh Douyin nguồn
- Cấu hình Workflow / Render / Voice / Google Drive
Dành cho dự án: Hệ thống AI Lấy Video TikTok / Douyin tự động.
"""

import os
import sys
import json
import csv
from pathlib import Path
from channel_output_paths import channel_review_path

# Đảm bảo UTF-8
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

BASE_DIR = Path("D:/0.1_FACE_BOOK/Lay video tiktok")
CONFIG_PATH = BASE_DIR / "backend" / "channels_config.json"
OUTPUT_XLSX = BASE_DIR / "DANH_SACH_KENH_DOUYIN_INPUT_OUTPUT.xlsx"
OUTPUT_CSV = BASE_DIR / "DANH_SACH_KENH_DOUYIN_INPUT_OUTPUT.csv"

# Thông tin mở rộng cho từng kênh (chủ đề Douyin, link mẫu, quy cách đầu vào/ra)
CHANNEL_METADATA = {
    "KENH_001_TienVaDoi": {
        "theme": "Triết lý, tài chính, đồng tiền & cuộc sống",
        "douyin_sample_link": "https://www.douyin.com/user/MS4wLjABAAAA... (Kênh Triết Lý / Tiền & Đời)",
        "input_folder": "inputs/KENH_001_TienVaDoi",
        "output_base": "outputs/KENH_001_TienVaDoi",
        "status": "Đang hoạt động",
        "note": "Black bar trên/dưới điện ảnh, voice Tài Chính trầm ấm"
    },
    "KENH_002_BatPhimLen": {
        "theme": "Review phim điện ảnh, phim ngắn kịch tính",
        "douyin_sample_link": "https://www.douyin.com/user/MS4wLjABAAAA... (Kênh Phim Điện Ảnh)",
        "input_folder": "inputs/KENH_002_BatPhimLen",
        "output_base": "outputs/KENH_002_BatPhimLen",
        "status": "Đang hoạt động",
        "note": "Title card 6s biến mất, sub vàng/trắng trên dải đen"
    },
    "KENH_003_ReviewSachCoTam": {
        "theme": "Tóm tắt sách hay, kỹ năng sống, bài học thành công",
        "douyin_sample_link": "https://www.douyin.com/user/MS4wLjABAAAA... (Kênh Sách / Trí Tuệ Douyin)",
        "input_folder": "inputs/KENH_003_ReviewSachCoTam",
        "output_base": "outputs/KENH_003_ReviewSachCoTam",
        "status": "Đang hoạt động",
        "note": "Clean sub, voice Better Version truyền cảm hứng"
    },
    "KENH_004_BoxmanKeChuyen": {
        "theme": "Kể chuyện kinh dị, trinh thám, vụ án bí ẩn",
        "douyin_sample_link": "https://www.douyin.com/user/MS4wLjABAAAA... (Kênh Kể Chuyện / Vụ Án Douyin)",
        "input_folder": "inputs/KENH_004_BoxmanKeChuyen",
        "output_base": "outputs/KENH_004_BoxmanKeChuyen",
        "status": "Đang hoạt động",
        "note": "Giọng Nam Minh 1.3x hào sảng, che đen đáy"
    },
    "KENH_005_VuTruCayPhim": {
        "theme": "Clone phim bộ nhiều tập, đa nhân vật nam nữ",
        "douyin_sample_link": "https://www.douyin.com/user/MS4wLjABAAAA... (Kênh Phim Bộ Ngắn Trung Quốc)",
        "input_folder": "inputs/KENH_005_VuTruCayPhim",
        "output_base": "outputs/KENH_005_VuTruCayPhim",
        "status": "Đang hoạt động",
        "note": "Clone phim đa nhân vật (F0 Whisper + Pyannote Diarization)"
    },
    "KENH_006_TieuWuby": {
        "theme": "Hoạt hình 2D/3D anime, hoạt họa vui nhộn, meme hài",
        "douyin_sample_link": "https://www.douyin.com/user/MS4wLjABAAAA... (Kênh Hoạt Hình Wuby)",
        "input_folder": "inputs/KENH_006",
        "output_base": "outputs/KENH_006",
        "status": "Đang hoạt động",
        "note": "Overlay trực tiếp lên video gốc, font hoạt hình, voice Wuby Clone cute"
    },
    "KENH_007_Song1MinhCoGiVui": {
        "theme": "Vlog đời sống độc thân, nấu ăn, chữa lành",
        "douyin_sample_link": "https://www.douyin.com/user/MS4wLjABAAAA... (Kênh Sống Một Mình Douyin)",
        "input_folder": "inputs/KENH_007",
        "output_base": "outputs/KENH_007",
        "status": "Đang hoạt động",
        "note": "Wuby overlay, voice Wuby Clone nhẹ nhàng 1.0x"
    },
    "KENH_008_ChiDepDiCau": {
        "theme": "Gái đẹp đi câu, dã ngoại ngoài trời, săn bắt dân dã",
        "douyin_sample_link": "https://www.douyin.com/user/MS4wLjABAAAA... (Kênh Gái Đẹp Câu Cá Douyin)",
        "input_folder": "inputs/KENH_008",
        "output_base": "outputs/KENH_008",
        "status": "Đang hoạt động",
        "note": "Làm mờ kính frosted glass che sub Trung, voice Nữ Cute 1.45x"
    },
    "-1004462729383": {
        "theme": "Nhóm Telegram Quản Trị & Test Render Toàn Hệ Thống",
        "douyin_sample_link": "Nhận link Douyin trực tiếp từ tin nhắn Telegram",
        "input_folder": "workspace (Tải tự động qua Bot)",
        "output_base": "workspace / outputs",
        "status": "Control Bot",
        "note": "Hỗ trợ test nhanh cả 4-5 cấu hình render và gửi trực tiếp về group"
    }
}


def build_data():
    with open(CONFIG_PATH, "r", encoding="utf-8") as f:
        config = json.load(f)

    channels = config.get("channels", {})
    voice_profiles = config.get("voice_profiles", {})

    headers = [
        "STT",
        "Mã Kênh Hệ Thống",
        "Tên Kênh Hiển Thị",
        "Chủ Đề Nội Dung Douyin",
        "Link Kênh Douyin Nguồn (Khai báo)",
        "Thư Mục Đầu Vào (Local Input)",
        "Thư Mục Đầu Ra Render (02_MOI_RENDER)",
        "Cấu Trúc 5 Thư Mục Vòng Đời",
        "Thư Mục Google Drive Đích",
        "Workflow",
        "Render Profile (Kiểu Che Sub)",
        "Giọng Đọc (Voice Profile)",
        "Tốc Độ (Speed)",
        "Logo / Watermark",
        "Trạng Thái",
        "Ghi Chú Kỹ Thuật"
    ]

    rows = []
    idx = 1
    for ch_key, ch_val in channels.items():
        name = ch_val.get("name", ch_key)
        workflow = ch_val.get("workflow", "Mặc định")
        render_profile = ch_val.get("render_profile", ch_val.get("mask_mode", "auto"))
        voice_key = ch_val.get("voice", "default")
        voice_info = voice_profiles.get(voice_key, {})
        voice_label = voice_info.get("label", voice_key)
        speed = f"{ch_val.get('speed', 1.0)}x"
        drive_folder = ch_val.get("drive_folder", "Chưa gán")
        
        logo_info = ch_val.get("logo", {})
        logo_tag = logo_info.get("tag", ch_val.get("_logo_type", "Chưa có"))

        meta = CHANNEL_METADATA.get(ch_key, {
            "theme": "Kênh tùy biến",
            "douyin_sample_link": "https://www.douyin.com/user/...",
            "input_folder": f"inputs/{ch_key}",
            "output_base": f"outputs/{ch_key}",
            "status": "Đang cấu hình",
            "note": ""
        })

        input_folder_full = f"D:/0.1_FACE_BOOK/Lay video tiktok/{meta['input_folder']}"
        output_render_full = f"D:/0.1_FACE_BOOK/Lay video tiktok/{meta['output_base']}/02_MOI_RENDER_CHO_XAC_NHAN"
        output_render_full = channel_review_path(BASE_DIR, ch_key, ch_val, output_render_full)
        lifecycle_folders = (
            f"01_CHUA_RENDER | 02_MOI_RENDER_CHO_XAC_NHAN | "
            f"03_CHUAN_BI_UP | 04_DA_UP | 99_FAILED"
        )

        rows.append([
            idx,
            ch_key,
            name,
            meta["theme"],
            meta["douyin_sample_link"],
            input_folder_full,
            output_render_full,
            lifecycle_folders,
            f"Google Drive: {drive_folder}",
            workflow,
            render_profile,
            voice_label,
            speed,
            logo_tag,
            meta["status"],
            meta["note"]
        ])
        idx += 1

    return headers, rows


def export_csv(headers, rows):
    # Ghi file CSV với utf-8-sig (BOM) để Microsoft Excel trên Windows mở trực tiếp không lỗi tiếng Việt
    with open(OUTPUT_CSV, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(headers)
        writer.writerows(rows)
    print(f"✅ Đã xuất CSV: {OUTPUT_CSV}")


def export_xlsx(headers, rows):
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "DANH_SACH_KENH_DOUYIN"

        # Tiêu đề bảng lớn
        ws.merge_cells("A1:P1")
        title_cell = ws["A1"]
        title_cell.value = "BẢNG QUẢN TRỊ THƯ MỤC ĐẦU VÀO - ĐẦU RA - LINK KÊNH DOUYIN DỰ ÁN VIDEO AI"
        title_cell.font = Font(name="Segoe UI", size=15, bold=True, color="FFFFFF")
        title_cell.fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 40

        # Subtitle
        ws.merge_cells("A2:P2")
        sub_cell = ws["A2"]
        sub_cell.value = "Hệ thống tự động: Link Douyin ➔ Tải về Input ➔ Tách Vocal/Diarize/OmniVoice ➔ Render Che Sub ➔ Output 5 Trạng Thái & Google Drive"
        sub_cell.font = Font(name="Segoe UI", size=10, italic=True, color="E0E6ED")
        sub_cell.fill = PatternFill(start_color="2B4C7E", end_color="2B4C7E", fill_type="solid")
        sub_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 24

        # Dòng Header
        header_fill = PatternFill(start_color="005A9C", end_color="005A9C", fill_type="solid")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        thin_border = Border(
            left=Side(style='thin', color='CCCCCC'),
            right=Side(style='thin', color='CCCCCC'),
            top=Side(style='thin', color='CCCCCC'),
            bottom=Side(style='thin', color='CCCCCC')
        )

        ws.row_dimensions[4].height = 32
        for col_idx, h_text in enumerate(headers, start=1):
            cell = ws.cell(row=4, column=col_idx)
            cell.value = h_text
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border

        # Dữ liệu
        row_alt_fill = PatternFill(start_color="F4F7FA", end_color="F4F7FA", fill_type="solid")
        data_font = Font(name="Segoe UI", size=10)

        for r_idx, row_data in enumerate(rows, start=5):
            ws.row_dimensions[r_idx].height = 28
            is_alt = (r_idx % 2 == 0)
            for c_idx, val in enumerate(row_data, start=1):
                cell = ws.cell(row=r_idx, column=c_idx)
                cell.value = val
                cell.font = data_font
                cell.border = thin_border
                
                # Căn lề
                if c_idx in [1, 13, 15]: # STT, Tốc độ, Trạng thái
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif c_idx in [2, 10, 11]: # Mã kênh, workflow
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

                if is_alt:
                    cell.fill = row_alt_fill

                # Màu sắc trạng thái
                if c_idx == 15 and val == "Đang hoạt động":
                    cell.font = Font(name="Segoe UI", size=10, bold=True, color="008000")
                elif c_idx == 15 and val == "Control Bot":
                    cell.font = Font(name="Segoe UI", size=10, bold=True, color="0055AA")

        # Tự động căn chỉnh độ rộng cột
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.row < 4:
                    continue
                if cell.value:
                    val_str = str(cell.value)
                    max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 48)

        # Tab thứ 2: HƯỚNG DẪN QUY TRÌNH & THƯ MỤC
        ws_guide = wb.create_sheet(title="HUONG_DAN_QUY_TRINH")
        ws_guide.column_dimensions["A"].width = 25
        ws_guide.column_dimensions["B"].width = 85

        guides = [
            ("QUY TRÌNH HỆ THỐNG", "MÔ TẢ CHI TIẾT"),
            ("1. Thư mục Đầu vào (Input)", "Chứa video Douyin gốc. Có thể tải bằng link kênh Douyin qua bot Telegram hoặc dán trực tiếp file .mp4 vào thư mục inputs/KENH_xxx/."),
            ("2. Vòng đời 5 thư mục Output", "Hệ thống tự động luân chuyển file qua 5 thư mục chuẩn của từng kênh:\n"
                                             "• 01_CHUA_RENDER: Video gốc vừa tải về, đang xếp hàng chờ chạy AI.\n"
                                             "• 02_MOI_RENDER_CHO_XAC_NHAN: Video đã che sub, dịch và lồng tiếng OmniVoice xong, sẵn sàng để xem thử/kiểm duyệt.\n"
                                             "• 03_CHUAN_BI_UP: Video đã được duyệt đạt chuẩn, chuẩn bị nạp vào tool đăng bài.\n"
                                             "• 04_DA_UP: Video đã đăng thành công lên TikTok / YouTube / Facebook.\n"
                                             "• 99_FAILED: Các video bị lỗi kỹ thuật hoặc bị từ chối."),
            ("3. Link Kênh Douyin", "Mỗi kênh trong hệ thống được cấu hình cho một ngách nội dung cụ thể (Phim, Sách, Hoạt hình, Câu cá,...). Bạn có thể thêm link kênh Douyin yêu thích vào cột 'Link Kênh Douyin Nguồn' để bot quét tự động."),
            ("4. Đồng bộ Google Sheet", "Dự án sử dụng Google Spreadsheet ID: 13U7eHfquutyBfrjtEtsqAFSXHrwxW9aqBSXjFnjUccs.\n"
                                        "Có thể đẩy bảng này trực tiếp lên Google Sheet thành tab mới bằng tools/push_channels_to_sheet.py."),
            ("5. TTS Engine Độc Tôn", "OmniVoice (kjanh/KhanhTTS-OmniVoice) chạy trên GPU CUDA FP16, bảo toàn chất giọng tự nhiên, cấm chỉnh tần số lệch chuẩn.")
        ]

        for r_i, (g_title, g_desc) in enumerate(guides, start=1):
            c1 = ws_guide.cell(row=r_i, column=1, value=g_title)
            c2 = ws_guide.cell(row=r_i, column=2, value=g_desc)
            if r_i == 1:
                c1.font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
                c1.fill = header_fill
                c2.font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
                c2.fill = header_fill
                ws_guide.row_dimensions[r_i].height = 28
            else:
                c1.font = Font(name="Segoe UI", size=10, bold=True)
                c2.font = Font(name="Segoe UI", size=10)
                c1.alignment = Alignment(vertical="top")
                c2.alignment = Alignment(vertical="top", wrap_text=True)
                ws_guide.row_dimensions[r_i].height = 45

        wb.save(OUTPUT_XLSX)
        print(f"✅ Đã xuất Excel: {OUTPUT_XLSX}")
    except Exception as e:
        print(f"❌ Lỗi xuất Excel: {e}")


def main():
    headers, rows = build_data()
    export_csv(headers, rows)
    export_xlsx(headers, rows)


if __name__ == "__main__":
    main()
